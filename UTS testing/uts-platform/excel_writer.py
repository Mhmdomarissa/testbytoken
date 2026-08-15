import os
from copy import copy
from openpyxl import load_workbook


class ExcelWriter:

    def __init__(self, template_file):
        self.template_file = template_file

    @staticmethod
    def _read_headers(ws, header_row=1):
        headers = {}
        for c in range(1, ws.max_column + 1):
            value = ws.cell(header_row, c).value
            if value:
                headers[str(value).strip()] = c
        return headers

    def _ensure_rownum_first_column(self, ws, header_row=1):
        headers = self._read_headers(ws, header_row)
        if "RowNum" not in headers:
            ws.insert_cols(1)
            ws.cell(header_row, 1).value = "RowNum"
            return self._read_headers(ws, header_row)

        if headers["RowNum"] == 1:
            return headers

        # Move RowNum to column 1 — read header names left-to-right, put RowNum first
        ordered = sorted(headers.items(), key=lambda item: item[1])
        names = ["RowNum"] + [name for name, _ in ordered if name != "RowNum"]
        for c in range(1, ws.max_column + 1):
            ws.cell(header_row, c).value = None
        for idx, name in enumerate(names, start=1):
            ws.cell(header_row, idx).value = name
        return self._read_headers(ws, header_row)

    def write(self, scenarios, output_file):

        if not os.path.exists(self.template_file):
            raise FileNotFoundError(self.template_file)

        wb = load_workbook(self.template_file)
        ws = wb.active

        # ----------------------------
        # Read Header Row
        # ----------------------------
        header_row = 1
        headers = self._ensure_rownum_first_column(ws, header_row)

        # ----------------------------
        # Collect Variables
        # ----------------------------
        variables = []

        for scenario in scenarios:

            vars = self.collect_variables(scenario)

            for key in vars.keys():
                if key not in variables:
                    variables.append(key)

        # ----------------------------
        # Create Missing Columns
        # ----------------------------
        col = ws.max_column + 1

        for var in variables:

            if var not in headers:
                ws.cell(header_row, col).value = var
                headers[var] = col
                col += 1

        # ----------------------------
        # Clear Existing Data
        # ----------------------------
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

        # ----------------------------
        # Write Data — RowNum count starts at 2 (row 1 is header only)
        # ----------------------------
        row = 2
        rownum_start = 2

        for offset, scenario in enumerate(scenarios):
            row_num = rownum_start + offset

            variables = self.collect_variables(scenario)

            if "RowNum" in headers:
                ws.cell(row, headers["RowNum"]).value = row_num

            if "AutomationTCName" in headers:
                ws.cell(
                    row,
                    headers["AutomationTCName"]
                ).value = scenario.get(
                    "name",
                    scenario.get("tc_id_name", "")
                )

            for key, value in variables.items():

                if key in headers:
                    ws.cell(
                        row,
                        headers[key]
                    ).value = value

            row += 1

        wb.save(output_file)

    # ----------------------------------------------------

    def collect_variables(self, scenario):

        variables = {}

        # Test Data
        test_data = scenario.get("test_data", "")

        if test_data:

            import re

            matches = re.findall(
                r"<([^>]+)>:([^<]+)",
                test_data
            )

            for name, value in matches:
                variables[f"<{name}>"] = value.strip()

        # Steps
        for block in scenario.get("functional_blocks", []):

            for step in block.get("steps", []):

                value = str(step.get("input_value", ""))

                if value.startswith("<") and ":" in value:

                    var = value.split(":")[0]
                    val = value.split(":", 1)[1]

                    variables[var] = val

                # Capture Expected Result

                if step.get("action") in (
                    "VerifyText",
                    "VerifyNonEditableText"
                ):

                    value = step.get("input_value", "")

                    if ":" in value:
                        variables["<ExpectedResult>"] = value.split(
                            ":",
                            1
                        )[1]
                    else:
                        variables["<ExpectedResult>"] = value

        return variables