"""Export discovery scenarios and run results to ALM-compatible formats."""

from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from automation.base import TestResult
from dynamic.discovery import DiscoveryResult, ScenarioDef

ROOT = Path(__file__).resolve().parent.parent
CONFIG_FILE = ROOT / "config" / "alm.json"

ALM_HEADERS = [
    "Test ID",
    "Test Name",
    "Module",
    "Priority",
    "Test Type",
    "Status",
    "Step #",
    "Step Description",
    "Action",
    "Input / Data",
    "Expected Result",
    "Automation Path",
]


def load_alm_config() -> dict:
    if CONFIG_FILE.is_file():
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    return {
        "output_path": "alm-output",
        "export_excel": True,
        "export_xml": True,
        "default_priority": "Medium",
        "default_test_type": "Automated",
        "default_status": "Design",
    }


def _scenario_rows(
    scenario: ScenarioDef,
    discovery: DiscoveryResult,
    cfg: dict,
    execution_status: str = "",
) -> list[list]:
    rows = []
    module = getattr(scenario, "module", "") or ""
    priority = cfg.get("default_priority", "Medium")
    test_type = cfg.get("default_test_type", "Automated")
    status = execution_status or cfg.get("default_status", "Design")
    automation_path = f"{scenario.id} / {discovery.app_name}"

    for step in scenario.steps:
        rows.append([
            scenario.id,
            scenario.title,
            module,
            priority,
            test_type,
            status,
            step.step_no,
            step.expected or step.object_name,
            step.action,
            step.input_value or "",
            step.expected or "",
            automation_path,
        ])
    return rows


def export_discovery_to_alm(
    discovery: DiscoveryResult,
    scenarios_filter: list[str] | None = None,
    execution_results: list[TestResult] | None = None,
) -> Path:
    """
    Export test cases to ALM Excel + XML.
    scenarios_filter: only export scenarios for these module names (Cycle 2).
    """
    cfg = load_alm_config()
    output_dir = ROOT / cfg.get("output_path", "alm-output")
    output_dir.mkdir(parents=True, exist_ok=True)

    result_map = {}
    if execution_results:
        result_map = {r.tc_id: r.status for r in execution_results}

    scenarios = list(discovery.scenarios)
    if scenarios_filter:
        wanted = {m.lower() for m in scenarios_filter}

        def _match(s: ScenarioDef) -> bool:
            mod = (getattr(s, "module", "") or "").strip().lower()
            if not mod or mod in ("login", "core"):
                return False
            return mod in wanted or any(w in mod for w in wanted)

        scenarios = [s for s in scenarios if _match(s)]

    all_rows: list[list] = []
    for scenario in scenarios:
        exec_status = result_map.get(scenario.id, "")
        if exec_status == "PASS":
            exec_status = "Passed"
        elif exec_status == "FAIL":
            exec_status = "Failed"
        all_rows.extend(_scenario_rows(scenario, discovery, cfg, exec_status))

    excel_path = output_dir / "ALM_TestCases.xlsx"
    xml_path = output_dir / "ALM_TestCases.xml"

    if cfg.get("export_excel", True):
        _write_alm_excel(excel_path, all_rows, discovery)

    if cfg.get("export_xml", True):
        _write_alm_xml(xml_path, scenarios, discovery, cfg, result_map)

    summary = {
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "app_name": discovery.app_name,
        "url": discovery.url,
        "test_case_count": len(scenarios),
        "step_count": len(all_rows),
        "excel": str(excel_path),
        "xml": str(xml_path),
    }
    (output_dir / "alm-export-summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    return output_dir


def export_modules_catalog_to_alm(
    modules: list[dict],
    app_name: str,
    url: str,
) -> Path:
    """Cycle 1: export scanned modules as ALM test case stubs."""
    cfg = load_alm_config()
    output_dir = ROOT / cfg.get("output_path", "alm-output")
    output_dir.mkdir(parents=True, exist_ok=True)

    rows = []
    for idx, mod in enumerate(modules, start=1):
        name = mod.get("name", f"Module_{idx}")
        tc_id = f"TC_ALM_{idx:03d}"
        rows.append([
            tc_id,
            f"{name} module functional test",
            name,
            cfg.get("default_priority", "Medium"),
            cfg.get("default_test_type", "Automated"),
            "Design",
            1,
            f"Login and open {name} module",
            "Trigger",
            url,
            f"{name} module page loads",
            f"{name} / {app_name}",
        ])

    excel_path = output_dir / "ALM_ModuleCatalog.xlsx"
    _write_alm_excel(excel_path, rows, app_name=app_name, url=url, title="Module Catalog")

    summary = {
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "app_name": app_name,
        "url": url,
        "module_count": len(modules),
        "excel": str(excel_path),
    }
    (output_dir / "alm-module-catalog.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    return output_dir


def _write_alm_excel(
    path: Path,
    rows: list[list],
    discovery: DiscoveryResult | None = None,
    app_name: str = "",
    url: str = "",
    title: str = "Test Cases",
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "ALM Test Cases"

    meta = [
        ("ALM Export", title),
        ("Application", discovery.app_name if discovery else app_name),
        ("URL", discovery.url if discovery else url),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for r, (label, value) in enumerate(meta, start=1):
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value)

    header_row = len(meta) + 2
    for c, header in enumerate(ALM_HEADERS, start=1):
        cell = ws.cell(header_row, c, header)
        cell.font = Font(bold=True)

    for r, row in enumerate(rows, start=header_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)

    wb.save(path)


def _write_alm_xml(
    path: Path,
    scenarios: list[ScenarioDef],
    discovery: DiscoveryResult,
    cfg: dict,
    result_map: dict[str, str],
) -> None:
    root = ET.Element("ALMTestSuite")
    root.set("app", discovery.app_name)
    root.set("url", discovery.url)
    root.set("generated", datetime.now().isoformat(timespec="seconds"))

    for scenario in scenarios:
        tc = ET.SubElement(root, "TestCase")
        ET.SubElement(tc, "TestID").text = scenario.id
        ET.SubElement(tc, "TestName").text = scenario.title
        ET.SubElement(tc, "Module").text = getattr(scenario, "module", "") or ""
        ET.SubElement(tc, "Priority").text = cfg.get("default_priority", "Medium")
        ET.SubElement(tc, "TestType").text = cfg.get("default_test_type", "Automated")
        status = result_map.get(scenario.id) or cfg.get("default_status", "Design")
        ET.SubElement(tc, "Status").text = status

        steps_el = ET.SubElement(tc, "Steps")
        for step in scenario.steps:
            st = ET.SubElement(steps_el, "Step")
            ET.SubElement(st, "StepNumber").text = str(step.step_no)
            ET.SubElement(st, "Description").text = step.object_name or ""
            ET.SubElement(st, "Action").text = step.action or ""
            ET.SubElement(st, "Input").text = step.input_value or ""
            ET.SubElement(st, "ExpectedResult").text = step.expected or ""

    tree = ET.ElementTree(root)
    ET.indent(tree, space="  ")
    tree.write(path, encoding="utf-8", xml_declaration=True)
