#version 7.0.0.38
import sys
import os
import glob
import json
import re
import csv
import zipfile
import tempfile
import base64
import time
import xml.etree.ElementTree as ET

def import_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font
        return Workbook, load_workbook, Font
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for offline Excel conversion. "
            "Install it with: pip install openpyxl"
        ) from exc


def import_docx():
    try:
        from docx import Document
        return Document
    except ImportError as exc:
        raise ImportError(
            "python-docx is required only for .docx uploads. "
            "Install it with: pip install python-docx"
        ) from exc

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

try:
    from config import (
        AI_PROVIDER,
        API_KEY,
        MODEL,
        AZURE_ENDPOINT,
        EXCEL_PATH,
        EXCEL_SHEET_NAME,
        TESTCASE_ROW,
        TC_ID_NAME_COLUMN,
        DATA_START_ROW_NO,
        TEST_DATA_COLUMN,
        TEST_STEPS_WHEN_COLUMN,
        EXPECTED_RESULT_THEN_COLUMN,
        OUTPUT_PATH as CONFIG_OUTPUT_PATH,
        TEMPLATE_SEARCH_PATHS as CONFIG_TEMPLATE_SEARCH_PATHS,
        WATCH_DEFAULT_FOLDER as CONFIG_WATCH_DEFAULT_FOLDER,
        LOG_FILE as CONFIG_LOG_FILE,
        PROMPT_TEXT as CONFIG_PROMPT_TEXT,
        DELIVERY_FOLDER as CONFIG_DELIVERY_FOLDER,
        INPUT_FOLDER,
        SCREENSHOTS_FOLDER,
        TEMPLATES_FOLDER,
        AUTO_FIND_EXCEL,
        SHARED_TESTDATA_FILE,
    )
except ImportError:
    AI_PROVIDER = "offline"
    API_KEY = ""
    MODEL = ""
    AZURE_ENDPOINT = ""
    EXCEL_PATH = ""
    EXCEL_SHEET_NAME = ""
    TESTCASE_ROW = 3
    TC_ID_NAME_COLUMN = 3
    DATA_START_ROW_NO = 2
    TEST_DATA_COLUMN = 8
    TEST_STEPS_WHEN_COLUMN = 9
    EXPECTED_RESULT_THEN_COLUMN = 10
    AUTO_FIND_EXCEL = True
    SHARED_TESTDATA_FILE = "TestData.xlsx"
    CONFIG_OUTPUT_PATH = os.path.join(SCRIPT_DIR, "output")
    INPUT_FOLDER = os.path.join(SCRIPT_DIR, "input")
    SCREENSHOTS_FOLDER = os.path.join(SCRIPT_DIR, "screenshots")
    TEMPLATES_FOLDER = os.path.join(SCRIPT_DIR, "templates")
    CONFIG_TEMPLATE_SEARCH_PATHS = [TEMPLATES_FOLDER]
    CONFIG_WATCH_DEFAULT_FOLDER = INPUT_FOLDER
    CONFIG_LOG_FILE = os.path.join(SCRIPT_DIR, "generator_log.txt")
    CONFIG_PROMPT_TEXT = (
        "Generate test cases from offline excel sheet. "
        "Final output should be test case for Xpedite."
    )
    CONFIG_DELIVERY_FOLDER = SCRIPT_DIR

if AI_PROVIDER == "openai" and API_KEY:
    from openai import OpenAI
    client = OpenAI(api_key=API_KEY)

elif AI_PROVIDER == "gemini" and API_KEY:
    from google import genai
    client = genai.Client(api_key=API_KEY)
    
elif AI_PROVIDER == "claude" and API_KEY:
    import anthropic
    client = anthropic.Anthropic(api_key=API_KEY)

elif AI_PROVIDER == "azure_openai" and API_KEY:
    from openai import AzureOpenAI
    client = AzureOpenAI(api_key=API_KEY,api_version="2024-10-21",azure_endpoint=AZURE_ENDPOINT)  
    
elif AI_PROVIDER == "ollama":
    from openai import OpenAI
    client = OpenAI(base_url=AZURE_ENDPOINT,api_key="ollama")

else:
    client = None

try:
    from PIL import Image
except ImportError:
    Image = None

import traceback
from datetime import datetime

from uts_engine.exporters.excel_writer import ExcelWriter

TEMPLATE_PATH = TEMPLATES_FOLDER
OUTPUT_PATH = CONFIG_OUTPUT_PATH
MAX_DOC_CHARS = 20000   # Prevent token overflow
WATCH_DEFAULT_FOLDER = CONFIG_WATCH_DEFAULT_FOLDER
EXCEL_EXTENSIONS = (".xlsx", ".xlsm", ".xltx", ".xltm")
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".bmp")
WATCH_EXTENSIONS = EXCEL_EXTENSIONS + (".csv", ".txt", ".md", ".json", ".docx", ".zip") + IMAGE_EXTENSIONS
TEMPLATE_SEARCH_PATHS = CONFIG_TEMPLATE_SEARCH_PATHS
DEFAULT_EXCEL_PATH = EXCEL_PATH
DEFAULT_PROMPT_TEXT = CONFIG_PROMPT_TEXT
DELIVERY_FOLDER = CONFIG_DELIVERY_FOLDER


def ensure_runtime_folders():
    for folder in (INPUT_FOLDER, SCREENSHOTS_FOLDER, OUTPUT_PATH, TEMPLATE_PATH):
        os.makedirs(folder, exist_ok=True)


def find_excel_in_folder(folder):
    if not folder or not os.path.isdir(folder):
        return None
    matches = []
    for name in os.listdir(folder):
        lowered = name.lower()
        if lowered.endswith(EXCEL_EXTENSIONS) and not name.startswith("~$"):
            matches.append(os.path.join(folder, name))
    if not matches:
        return None
    matches.sort(key=os.path.getmtime, reverse=True)
    return matches[0]


def resolve_excel_path(file_path=None):
    if file_path and os.path.exists(file_path):
        return file_path
    if EXCEL_PATH and os.path.exists(EXCEL_PATH):
        return EXCEL_PATH
    if AUTO_FIND_EXCEL:
        found = find_excel_in_folder(INPUT_FOLDER)
        if found:
            log(f"Auto-detected Excel: {found}")
            return found
    return file_path or EXCEL_PATH or None


def get_screenshots_folder():
    if os.path.isdir(SCREENSHOTS_FOLDER):
        return SCREENSHOTS_FOLDER
    return None


def folder_has_images(folder):
    if not folder or not os.path.isdir(folder):
        return False
    return any(
        name.lower().endswith(IMAGE_EXTENSIONS)
        for name in os.listdir(folder)
    )


def resolve_template_path(filename):
    for folder in TEMPLATE_SEARCH_PATHS:
        candidate = os.path.join(folder, filename)
        if os.path.exists(candidate):
            return candidate
    raise FileNotFoundError(
        f"Template not found: {filename}. "
        f"Checked: {', '.join(TEMPLATE_SEARCH_PATHS)}"
    )


def append_manual_testcase(
        ws,
        tc_id,
        scenario_name,
        bc_name,
        step_no,
        step):

    ws.append([
        tc_id,
        scenario_name,
        bc_name,
        step_no,
        step.get("object_name", ""),
        step.get("action", ""),
        step.get("input_value", ""),
        step.get("expected_value", ""),
        step.get("description", "")
    ])
    
def generate_manual_excel(scenarios):

    Workbook, _, Font = import_openpyxl()
    os.makedirs(OUTPUT_PATH, exist_ok=True)

    excel_file = os.path.join(
        OUTPUT_PATH,
        "Manual_Test_Cases.xlsx"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Manual Test Cases"

    headers = [
        "Test Case ID",
        "Scenario",
        "Step No",
        "Input Test Data",
        "Test Step Description",
        "Expected Result"
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)

    row_no = 2

    for tc_index, scenario in enumerate(scenarios, start=1):

        tc_id = f"TC_{tc_index:03d}"

        step_no = 1

        for block in scenario.get("functional_blocks", []):

            bc_name = block.get("bc_name", "")

            for step in block.get("steps", []):

                ws.cell(row_no,1).value = tc_id
                ws.cell(row_no,2).value = scenario.get("name","")
                ws.cell(row_no,3).value = step_no
                ws.cell(row_no,4).value = step.get("input_value","")
                ws.cell(row_no,5).value = step.get("description","")      
                if "verify" in step.get("description","").lower():
                    ws.cell(row_no,6).value = step.get("description","")
                    ws.cell(row_no,5).value = ""
                    
                row_no += 1
                step_no += 1

    wb.save(excel_file)


def call_ai(system_prompt, combined_text, image_folder=None):

    if AI_PROVIDER == "openai":

        user_content = [
            {
                "type": "text",
                "text": combined_text
            }
        ]

        # Attach all extracted images
        if image_folder and os.path.exists(image_folder):

            image_files = [
                f for f in os.listdir(image_folder)
                if f.lower().endswith((".png", ".jpg", ".jpeg"))
            ]

            for file in image_files:

                image_path = os.path.join(image_folder, file)

                with open(image_path, "rb") as f:
                    base64_image = base64.b64encode(
                        f.read()
                    ).decode("utf-8")

                user_content.append(
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/png;base64,{base64_image}"
                        }
                    }
                )

        messages = [
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_content
            }
        ]

        response = client.chat.completions.create(
            model=MODEL,
            temperature=0,
            max_tokens=4000,
            messages=messages
        )

        return response.choices[0].message.content.strip()

    elif AI_PROVIDER == "gemini":

        contents = [f"{system_prompt}\n\n{combined_text}"]

        if image_folder and os.path.exists(image_folder):

            image_files = [
                f for f in os.listdir(image_folder)
                if f.lower().endswith((".png", ".jpg", ".jpeg"))
            ]

            for file in image_files:
                try:
                    img = Image.open(os.path.join(image_folder, file))
                    contents.append(img)
                except Exception:
                    pass

        response = client.models.generate_content(
            model=MODEL,
            contents=contents
        )

        return response.text.strip()
        
    elif AI_PROVIDER == "claude":

        response = client.messages.create(
            model=MODEL,
            max_tokens=4000,
            temperature=0,
            system=system_prompt,
            messages=[
                {
                    "role": "user",
                    "content": combined_text
                }
            ]
        )

        return response.content[0].text
        
    elif AI_PROVIDER == "azure_openai":

        response = client.chat.completions.create(
        model=MODEL,  # Azure Deployment Name
        temperature=0,
        max_tokens=4000,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": combined_text
            }
        ]
        )

        return response.choices[0].message.content.strip()
        
    elif AI_PROVIDER == "ollama":

        response = client.chat.completions.create(
        model=MODEL,  
        temperature=0,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": combined_text
            }
        ]
        )

        return response.choices[0].message.content.strip()

    else:
        raise Exception(f"Unsupported provider: {AI_PROVIDER}")

# -------------------------------------------------
# LOG FUNCTION
# -------------------------------------------------
LOG_FILE = CONFIG_LOG_FILE

def log(message):
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{datetime.now()} - {message}\n")
        
        
def extract_static_variables(scenario):

    variables = {}

    for block in scenario.get("functional_blocks", []):
        for step in block.get("steps", []):

            input_val = step.get("input_value", "")

            # Match <Variable>:value (value may contain spaces)
            for var, val in re.findall(r"<([^>]+)>:([^<]+)", input_val):
                variables[f"<{var}>"] = val.strip()

    return variables

# -------------------------------------------------
# Utility
# -------------------------------------------------

def _business_flow_bc_name(component="", description="", scenario_name="", fallback="Business_Flow"):
    """Group related Excel steps under one BC for the business process."""
    try:
        from uts_engine.discovery.module_filter import business_flow_bc_name as _bf
        return _bf(component, description, scenario_name, fallback)
    except Exception:
        raw = component or description or scenario_name or fallback
        clean = re.sub(r"[^a-zA-Z0-9_ ]", "", str(raw).replace("-", "_"))
        return (clean.strip().replace(" ", "_") or fallback)[:60]


def sanitize_name(name):
    clean = re.sub(r'[^a-zA-Z0-9_ ]', '', str(name))
    return clean.strip().replace(" ", "_")


def format_tc_output_name(tc_id_name):
    clean = sanitize_name(tc_id_name)
    if not clean:
        return "TC_UnnamedScenario"
    if clean.upper().startswith("TC_"):
        return clean
    return f"TC_{clean}"


def tc_id_to_base(tc_id_name):
    base = sanitize_name(tc_id_name)
    if not base:
        return "UnnamedScenario"
    if base.upper().startswith("TC_"):
        return base[3:] or base
    return base


def resolve_output_names(scenario):
    tc_id_name = (
        scenario.get("tc_id_name")
        or scenario.get("tc_id")
        or scenario.get("name", "UnnamedScenario")
    )
    tc_name = format_tc_output_name(tc_id_name)
    base_name = tc_id_to_base(tc_id_name)
    bpw_name = f"BPW_{base_name}"
    return base_name, tc_name, bpw_name


def extract_requested_count(prompt, doc_text=None):

    # Explicit numeric request
    match = re.search(
        r'(\d+)\s*(test case|test cases|scenario|scenarios)',
        prompt,
        re.IGNORECASE
    )

    if match:
        return int(match.group(1))

    # Detect from uploaded document
    if doc_text:

        tc_matches = re.findall(
            r'Test Case\s*:\s*\d+',
            doc_text,
            re.IGNORECASE
        )

        if tc_matches:
            return len(tc_matches)

    # If prompt asks for ALL
    if re.search(r'\ball\b', prompt, re.IGNORECASE):
        return 20

    return 10


def clean_json_response(text):
    text = text.replace("```json", "").replace("```", "").strip()
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1:
        return text[start:end + 1]
    return text


def build_default_prompt(file_path=None):
    base_prompt = (
        "Generate test cases from offline excel sheet. "
        "Final output should be test case for Xpedite."
    )
    if file_path and re.search(r"orangehrm|hrm", os.path.basename(file_path), re.IGNORECASE):
        base_prompt += (
            " Generate detailed business workflow including create/update/delete/verify "
            "for https://opensource-demo.orangehrmlive.com/. "
            "Login with Admin/admin123."
        )
    return base_prompt


def encode_image(image_path):
    with open(image_path, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


# -------------------------------------------------
# Action Normalizer
# -------------------------------------------------
def normalize_action(action):
    allowed = [
        "InvokeURL",
        "InvokeApplication",
        "SetText",
        "SetPassword",
        "PerformSelect",
        "PerformClick",
        "VerifyText",
        "VerifyNonEditableText",
        "GetDataFrmFile",
        "GetExcelRowValues",
    ]
    if action in allowed:
        return action

    fallback = {
        "click": "PerformClick",
        "verify": "VerifyNonEditableText",
        "launch": "InvokeURL",
        "open": "InvokeURL",
        "type": "SetText",
        "select": "PerformSelect",
        "password": "SetPassword",
        "invokeapplication": "InvokeApplication",
        "mobile": "InvokeApplication"
    }
    return fallback.get(str(action).lower(), "PerformClick")


# -------------------------------------------------
# Generic Input Value Normalizer (Random Value)
# -------------------------------------------------
def generate_random_value():
    return "SampleValue"


def normalize_input_value(value):

    if not value:
        return value

    value = str(value).strip()

    # If already contains default value → keep as is
    if re.match(r"<[^>]+>:.+", value):
        return value

    # Detect pattern <VariableName>
    match = re.fullmatch(r"<([^>]+)>", value)

    if match:
        var_name = match.group(1)
        random_value = generate_random_value()
        return f"<{var_name}>:{random_value}"

    return value


def resolve_variable_default(var_name, test_data=""):
    variables = parse_test_data_variables(test_data)
    key = f"<{var_name}>"
    if key in variables:
        return variables[key]

    fallbacks = {
        "Username": "test_user",
        "Password": "Pass@123",
        "ExpectedResult": "SampleValue",
        "VerifyText": "SampleValue",
    }
    return fallbacks.get(var_name, generate_random_value())


def parameterize_verify_input_value(value, test_data="", default_var_name="ExpectedResult"):
    value = str(value or "").strip()
    if not value:
        return f"<{default_var_name}>:{generate_random_value()}"

    if re.match(r"<[^>]+>:.+", value):
        return value

    bare_var = re.fullmatch(r"<([^>]+)>", value)
    if bare_var:
        var_name = bare_var.group(1)
        return f"<{var_name}>:{resolve_variable_default(var_name, test_data)}"

    variables = parse_test_data_variables(test_data)
    for var_key, var_val in variables.items():
        if str(var_val).strip() == value:
            return f"{var_key}:{value}"

    return f"<{default_var_name}>:{value}"


def finalize_step(step, test_data=""):
    action = normalize_action(step.get("action", ""))
    step["action"] = action

    if action in ("VerifyNonEditableText", "VerifyText"):
        step["action"] = "VerifyNonEditableText"
        step["object_name"] = step.get("object_name") or "NonEditableText"

        input_value = str(step.get("input_value", "")).strip()
        expected_value = str(step.get("expected_value", "")).strip()
        verify_target = input_value

        if expected_value and expected_value.lower() != "true":
            if not input_value or input_value.lower() == "true":
                verify_target = expected_value
            elif not verify_target:
                verify_target = expected_value

        if not verify_target:
            verify_target = step.get("description", "")

        step["input_value"] = parameterize_verify_input_value(
            verify_target,
            test_data,
        )
        step["expected_value"] = "True"
        return step

    if action == "PerformClick":
        step["input_value"] = ""
        step["expected_value"] = ""
        return step

    input_value = step.get("input_value", "")
    if input_value:
        step["input_value"] = normalize_input_value(input_value)

    return step


def normalize_scenario_steps(scenario):
    test_data = scenario.get("test_data", "")
    for block in scenario.get("functional_blocks", []):
        block["steps"] = [
            finalize_step(step, test_data)
            for step in block.get("steps", [])
        ]
    return scenario

# -------------------------------------------------
# File Extraction
# -------------------------------------------------
def extract_zip(zip_path):
    temp_dir = tempfile.mkdtemp()
    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
        zip_ref.extractall(temp_dir)
    return temp_dir


def extract_docx_content(docx_path):
    Document = import_docx()
    temp_dir = tempfile.mkdtemp()
    doc = Document(docx_path)

    full_text = ""
    for para in doc.paragraphs:
        full_text += para.text + "\n"

    # Trim large documents
    if len(full_text) > MAX_DOC_CHARS:
        full_text = full_text[:MAX_DOC_CHARS]

    # Extract images
    with zipfile.ZipFile(docx_path, 'r') as docx:
        for file in docx.namelist():
            if file.startswith("word/media/"):
                docx.extract(file, temp_dir)

    media_folder = os.path.join(temp_dir, "word", "media")

    return full_text, media_folder


def extract_text_file_content(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()[:MAX_DOC_CHARS]


def extract_csv_content(csv_path):
    rows = []
    with open(csv_path, "r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for row_index, row in enumerate(reader, start=1):
            if row_index > 200:
                break
            cleaned = [str(cell).strip() for cell in row if str(cell).strip()]
            if cleaned:
                rows.append(" | ".join(cleaned))
    return "\n".join(rows)[:MAX_DOC_CHARS]


def extract_excel_content(excel_path):
    _, load_workbook, _ = import_openpyxl()
    if excel_path.lower().endswith(".xls"):
        raise ValueError(
            "Legacy .xls files are not supported directly. "
            "Please save the spreadsheet as .xlsx and upload again."
        )

    wb = load_workbook(excel_path, data_only=True, read_only=True)
    sections = []

    for sheet in wb.worksheets[:10]:
        sections.append(f"[Sheet: {sheet.title}]")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_index > 200:
                sections.append("... truncated ...")
                break
            values = []
            for cell in row[:20]:
                if cell is None:
                    continue
                cell_text = str(cell).strip()
                if cell_text:
                    values.append(cell_text)
            if values:
                sections.append(" | ".join(values))

    return "\n".join(sections).strip()[:MAX_DOC_CHARS]


def load_input_content(file_path):
    doc_text = None
    image_folder = None
    lowered = file_path.lower()

    if lowered.endswith(".zip"):
        log(" ZIP detected → Vision mode")
        log("ZIP detected")
        image_folder = extract_zip(file_path)

    elif lowered.endswith(".docx"):
        log(" Word detected → Full document + Vision mode")
        log("DOCX detected")
        doc_text, image_folder = extract_docx_content(file_path)

    elif lowered.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        log(" Excel detected → Spreadsheet parsing mode")
        doc_text = extract_excel_content(file_path)

    elif lowered.endswith(".csv"):
        log(" CSV detected → Text parsing mode")
        doc_text = extract_csv_content(file_path)

    elif lowered.endswith((".txt", ".md", ".json")):
        log(" Text document detected → Text parsing mode")
        doc_text = extract_text_file_content(file_path)

    elif lowered.endswith(".xls"):
        raise ValueError(
            "Legacy .xls files are not supported directly. "
            "Please convert the file to .xlsx and upload again."
        )

    return doc_text, image_folder


# -------------------------------------------------
# Offline Manual Excel -> Xpedite
# -------------------------------------------------
MANUAL_TC_HEADER_ALIASES = {
    "tc_id": ("tc id", "test case id", "test case #", "tcid"),
    "tc_id_name": ("tc id", "tc id name", "test case name", "tc name", "scenario name"),
    "uj_id": ("uj id", "user journey id"),
    "scenario": (
        "test scenario",
        "scenario",
        "test title",
        "sub module",
    ),
    "description": (
        "test case description",
        "test description",
        "precondation",
        "precondition",
        "description",
    ),
    "steps": ("test steps", "steps"),
    "expected": ("expected result", "expected"),
    "test_data": ("test data", "testdata"),
    "component": ("component", "module name", "system name"),
    "automate": (
        "automate",
        "automation feasibility",
        "automation",
        "test type",
    ),
}


def normalize_header(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def find_header_row(worksheet, max_rows=10):
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_rows, values_only=True),
        start=1,
    ):
        headers = [normalize_header(cell) for cell in row]
        has_steps = any(
            any(alias in header for alias in MANUAL_TC_HEADER_ALIASES["steps"])
            for header in headers
        )
        has_tc_id = any(
            any(alias in header for alias in MANUAL_TC_HEADER_ALIASES["tc_id"])
            for header in headers
        )
        if has_steps and has_tc_id:
            return row_index, headers
    return None, []


def map_manual_columns(headers):
    mapping = {}
    for key, aliases in MANUAL_TC_HEADER_ALIASES.items():
        for index, header in enumerate(headers):
            if any(alias in header for alias in aliases):
                mapping[key] = index
                break
    return mapping


def config_enabled_offline_mapping():
    return bool(EXCEL_SHEET_NAME) and all(
        value and int(value) > 0
        for value in (
            TESTCASE_ROW,
            TC_ID_NAME_COLUMN,
            DATA_START_ROW_NO,
            TEST_DATA_COLUMN,
            TEST_STEPS_WHEN_COLUMN,
            EXPECTED_RESULT_THEN_COLUMN,
        )
    )


def get_configured_sheet(workbook):
    if EXCEL_SHEET_NAME:
        if EXCEL_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Configured sheet '{EXCEL_SHEET_NAME}' not found. "
                f"Available sheets: {', '.join(workbook.sheetnames)}"
            )
        return workbook[EXCEL_SHEET_NAME]
    return workbook.worksheets[0]


def row_has_content(row_values):
    return any(cell is not None and str(cell).strip() for cell in row_values)


def pad_row(row_values, target_index):
    row_list = list(row_values)
    if len(row_list) <= target_index:
        row_list.extend([None] * (target_index + 1 - len(row_list)))
    return row_list


def is_manual_testcase_workbook(excel_path):
    _, load_workbook, _ = import_openpyxl()
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        worksheet = get_configured_sheet(wb)
        if config_enabled_offline_mapping():
            return True
        _, headers = find_header_row(worksheet)
        columns = map_manual_columns(headers)
        return "steps" in columns and "tc_id" in columns
    finally:
        wb.close()


def parse_numbered_steps(test_steps_text):
    text = cell_text(test_steps_text)
    if not text:
        return []

    # Handle multi-line and single-line numbered steps:
    # "1. Login", "1.Open app", "2) Click"
    parts = re.split(r"(?=\s*\d+[\).\:-]\s*)", text)
    if len(parts) <= 1:
        parts = re.split(r"(?=\s*\d+[\).\:-](?=\S))", text)

    steps = []
    for part in parts:
        cleaned = re.sub(r"^\s*\d+[\).\:-]\s*", "", part.strip())
        if not cleaned:
            cleaned = re.sub(r"^\s*\d+[\).\:-](?=\S)", "", part.strip())
        if cleaned:
            steps.append(cleaned)

    if not steps and text:
        steps.append(text)

    return steps


def extract_quoted_values(text):
    return re.findall(r'"([^"]+)"|\'([^\']+)\'', text)


def first_quoted_text(text):
    for match in extract_quoted_values(text):
        value = match[0] or match[1]
        if value:
            return value
    return ""


def extract_url(text):
    match = re.search(r"https?://[^\s\"']+", text or "", re.IGNORECASE)
    return match.group(0) if match else ""


def parse_test_data_variables(test_data):
    variables = {}
    text = cell_text(test_data)
    if not text:
        return variables

    for label, value in re.findall(
        r"([A-Za-z][A-Za-z0-9 _]*)\s*[:=]\s*([^\n,;|]+)",
        text,
    ):
        key = label.strip().replace(" ", "")
        val = value.strip()
        if key and val:
            variables[f"<{key}>"] = val
            if key.lower() == "username":
                variables["<Username>"] = val
            if key.lower() == "password":
                variables["<Password>"] = val

    lowered = text.lower()
    if "username" in lowered:
        variables.setdefault("<Username>", "test_user")
    if "password" in lowered:
        variables.setdefault("<Password>", "Pass@123")
    if "login" in lowered or "credential" in lowered:
        variables.setdefault("<Username>", "test_user")
        variables.setdefault("<Password>", "Pass@123")

    account_match = re.search(r"account\s*[:=]?\s*([A-Za-z0-9_-]+)", text, re.IGNORECASE)
    if account_match:
        variables["<FromAccount>"] = account_match.group(1)

    amount_match = re.search(r"amount\s*[:=]?\s*([0-9.,]+)", text, re.IGNORECASE)
    if amount_match:
        variables["<Amount>"] = amount_match.group(1)

    return variables


def make_step(object_name, action, input_value="", expected_value="", description=""):
    return {
        "object_name": object_name,
        "action": action,
        "input_value": input_value,
        "expected_value": expected_value,
        "description": description,
    }


def expand_login_steps(step_text, test_data):
    variables = parse_test_data_variables(test_data)
    username = variables.get("<Username>", "test_user")
    password = variables.get("<Password>", "Pass@123")
    steps = []
    if re.search(r"\b(open|launch|start)\b", step_text, re.IGNORECASE):
        steps.append(
            make_step(
                infer_screen_block_name(step_text),
                "PerformClick",
                "",
                "",
                step_text,
            )
        )
    if re.search(r"\b(click|press)\b", step_text, re.IGNORECASE):
        object_name = first_quoted_text(step_text) or "Login"
        steps.append(
            make_step(object_name, "PerformClick", "", "", step_text)
        )
    steps.extend([
        make_step(
            "Username",
            "SetText",
            f"<Username>:{username}",
            "",
            step_text,
        ),
        make_step(
            "Password",
            "SetPassword",
            f"<Password>:{password}",
            "",
            "Enter password",
        ),
        make_step("Login", "PerformClick", "", "", "Click Login"),
    ])
    return steps


def infer_screen_block_name(step_text):
    quoted = first_quoted_text(step_text)
    if quoted:
        return sanitize_name(quoted)

    navigate_match = re.search(
        r"navigate to (?:the )?(.+?)(?: tab| screen| page| menu|$)",
        step_text,
        re.IGNORECASE,
    )
    if navigate_match:
        return sanitize_name(navigate_match.group(1))

    click_match = re.search(
        r"click (?:on )?(?:the )?(.+?)(?: button| link| tab|$)",
        step_text,
        re.IGNORECASE,
    )
    if click_match:
        return sanitize_name(click_match.group(1))

    return sanitize_name(step_text[:40] or "Screen")


def infer_business_block_name(step_text, fallback="Business_Flow"):
    """Prefer explicit screen/page names so BCs stay screen-wise."""
    lowered = (step_text or "").lower()

    opened_match = re.search(
        r"opened ['\"]?(.+?)['\"]? (?:module|sub-page|screen|page|tab)",
        step_text,
        re.IGNORECASE,
    )
    if opened_match:
        return sanitize_name(opened_match.group(1))

    if re.search(r"\b(log in to|login to|sign in to|log into)\b", lowered):
        return "Login"

    inferred = infer_screen_block_name(step_text)
    if inferred and inferred.lower() not in ("verify_result", "screen"):
        return inferred
    return sanitize_name(fallback) or "Business_Flow"


def convert_manual_step_to_xpedite(step_text, test_data=""):
    lowered = step_text.lower()

    if re.search(r"\b(verify|validate|confirm|check)\b", lowered):
        target = first_quoted_text(step_text) or step_text
        return [
            make_step(
                "NonEditableText",
                "VerifyNonEditableText",
                parameterize_verify_input_value(target, test_data),
                "True",
                step_text,
            )
        ], "Verify_Result"

    if re.search(r"\b(select|choose|pick)\b", lowered):
        object_name = first_quoted_text(step_text) or infer_screen_block_name(step_text)
        var_name = sanitize_name(object_name).replace("_", "")
        return [
            make_step(
                object_name,
                "PerformSelect",
                f"<{var_name}>:{object_name}",
                "",
                step_text,
            )
        ], infer_screen_block_name(step_text)

    if re.search(r"\b(enter|type|input|fill|provide)\b", lowered):
        object_name = first_quoted_text(step_text)
        if not object_name:
            field_match = re.search(
                r"(?:enter|type|input|fill|provide)\s+(.+?)(?:\s+as|\s+with|\s+in|$)",
                step_text,
                re.IGNORECASE,
            )
            object_name = field_match.group(1).strip() if field_match else "Input Field"
        var_name = sanitize_name(object_name).replace("_", "") or "InputValue"
        variables = parse_test_data_variables(test_data)
        if "username" in lowered:
            default_value = variables.get("<Username>", object_name)
        elif "password" in lowered:
            default_value = variables.get("<Password>", object_name)
        else:
            default_value = variables.get(f"<{var_name}>", object_name)
        action = "SetPassword" if "password" in lowered else "SetText"
        return [
            make_step(
                object_name,
                action,
                f"<{var_name}>:{default_value}",
                "",
                step_text,
            )
        ], infer_screen_block_name(step_text)

    if re.search(r"\b(click|press|navigate|open|go to|access)\b", lowered):
        object_name = first_quoted_text(step_text) or infer_screen_block_name(step_text)
        return [
            make_step(object_name, "PerformClick", "", "", step_text)
        ], infer_screen_block_name(step_text)

    if re.search(r"\b(log in to|login to|sign in to|log into)\b", lowered):
        return expand_login_steps(step_text, test_data), "Login"

    return [
        make_step(infer_screen_block_name(step_text), "PerformClick", "", "", step_text)
    ], infer_screen_block_name(step_text)


def build_scenario_from_manual_row(row, columns, sheet_name=""):
    tc_id = cell_text(row[columns["tc_id"]]) if "tc_id" in columns else ""
    tc_id_name = cell_text(row[columns["tc_id_name"]]) if "tc_id_name" in columns else ""
    scenario_name = cell_text(row[columns["scenario"]]) if "scenario" in columns else tc_id
    if not tc_id_name:
        tc_id_name = scenario_name or tc_id
    description = cell_text(row[columns["description"]]) if "description" in columns else ""
    expected_result = cell_text(row[columns["expected"]]) if "expected" in columns else ""
    test_data = cell_text(row[columns["test_data"]]) if "test_data" in columns else ""
    component = cell_text(row[columns["component"]]) if "component" in columns else ""
    steps_text = row[columns["steps"]] if "steps" in columns else ""

    if "automate" in columns:
        automate_flag = cell_text(row[columns["automate"]]).lower()
        if automate_flag and automate_flag not in (
            "automate",
            "automation",
            "yes",
            "y",
        ):
            return None

    manual_steps = parse_numbered_steps(steps_text)
    if not manual_steps:
        return None

    name = scenario_name or tc_id or "Unnamed_Scenario"
    if description and description != name:
        name = f"{name}_{description}" if scenario_name else description

    tc_prefix = sanitize_name(sheet_name) if sheet_name else ""
    if tc_id_name:
        scenario_key = format_tc_output_name(tc_id_name)
    elif tc_id:
        scenario_key = format_tc_output_name(tc_id)
    elif tc_prefix:
        scenario_key = sanitize_name(f"{tc_prefix}_{name}")
    else:
        scenario_key = sanitize_name(name)
    functional_blocks = []
    all_text = " ".join([description, test_data, expected_result, " ".join(manual_steps)])
    app_url = extract_url(all_text)

    if app_url:
        functional_blocks.append({
            "bc_name": "Launch_Application",
            "steps": [
                make_step(
                    "ApplicationURL",
                    "InvokeURL",
                    f"<ApplicationURL>:{app_url}",
                    "",
                    "Launch application URL",
                )
            ],
        })

    # Keep BCs screen-wise: Launch + Login remain protected, while
    # business steps are grouped by inferred page/screen transitions.
    flow_bc_name = _business_flow_bc_name(
        component=component,
        description=description,
        scenario_name=scenario_name or tc_id or "Business_Flow",
    )

    current_block_name = None
    current_steps = []

    def flush_block():
        nonlocal current_block_name, current_steps
        if current_steps:
            functional_blocks.append({
                "bc_name": current_block_name or flow_bc_name or "Business_Flow",
                "steps": current_steps,
            })
            current_steps = []

    for step_text in manual_steps:
        converted_steps, inferred = convert_manual_step_to_xpedite(step_text, test_data)
        inferred_l = (inferred or "").lower()
        if inferred_l == "login" or inferred_l in ("username", "password"):
            target = "Login"
        elif inferred_l == "launch_application" or "launch" in inferred_l:
            target = "Launch_Application"
        else:
            target = infer_business_block_name(step_text, flow_bc_name)
        if target != current_block_name:
            flush_block()
            current_block_name = target
        current_steps.extend(converted_steps)

    flush_block()

    if expected_result and not re.search(r"\bverify\b", " ".join(manual_steps), re.IGNORECASE):
        functional_blocks.append({
            "bc_name": "Verify_Result",
            "steps": [
                make_step(
                    "NonEditableText",
                    "VerifyNonEditableText",
                    parameterize_verify_input_value(expected_result, test_data),
                    "True",
                    expected_result,
                )
            ],
        })

    scenario = {
        "name": scenario_key,
        "display_name": tc_id_name or name,
        "tc_id": tc_id,
        "tc_id_name": tc_id_name,
        "test_data": test_data,
        "component": component,
        "description": description,
        "functional_blocks": functional_blocks,
    }
    return merge_functional_blocks(scenario)


def build_scenario_from_config_row(row, row_number, sheet_name=""):
    test_case_col = max(0, int(TESTCASE_ROW) - 1)
    tc_id_name_col = max(0, int(TC_ID_NAME_COLUMN) - 1)
    test_data_col = max(0, int(TEST_DATA_COLUMN) - 1)
    test_steps_col = max(0, int(TEST_STEPS_WHEN_COLUMN) - 1)
    expected_col = max(0, int(EXPECTED_RESULT_THEN_COLUMN) - 1)
    max_col = max(test_case_col, tc_id_name_col, test_data_col, test_steps_col, expected_col)
    row = pad_row(row, max_col)

    tc_id = cell_text(row[test_case_col]) or f"TC_Row_{row_number}"
    tc_id_name = cell_text(row[tc_id_name_col]) or tc_id
    test_data = cell_text(row[test_data_col])
    steps_text = cell_text(row[test_steps_col])
    expected_result = cell_text(row[expected_col])

    if not steps_text:
        return None

    manual_steps = parse_numbered_steps(steps_text)
    if not manual_steps:
        return None

    columns = {
        "tc_id": test_case_col,
        "scenario": test_case_col,
        "description": test_case_col,
        "expected": expected_col,
        "test_data": test_data_col,
        "steps": test_steps_col,
    }
    scenario = build_scenario_from_manual_row(row, columns, sheet_name=sheet_name)
    if not scenario:
        return None

    scenario["tc_id"] = tc_id
    scenario["tc_id_name"] = tc_id_name
    scenario["display_name"] = tc_id_name
    scenario["name"] = format_tc_output_name(tc_id_name)
    scenario["description"] = tc_id_name
    scenario["test_data"] = test_data
    normalize_scenario_steps(scenario)
    return scenario


def parse_offline_manual_excel(excel_path):
    _, load_workbook, _ = import_openpyxl()
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    scenarios = []
    parsed_any_sheet = False

    try:
        worksheets = [get_configured_sheet(wb)] if EXCEL_SHEET_NAME else wb.worksheets
        for worksheet in worksheets:
            if config_enabled_offline_mapping():
                parsed_any_sheet = True
                sheet_name = worksheet.title
                for row_number, row in enumerate(
                    worksheet.iter_rows(
                        min_row=int(DATA_START_ROW_NO),
                        values_only=True,
                    ),
                    start=int(DATA_START_ROW_NO),
                ):
                    if not row or not row_has_content(row):
                        continue
                    scenario = build_scenario_from_config_row(
                        row,
                        row_number,
                        sheet_name=sheet_name,
                    )
                    if scenario and scenario.get("functional_blocks"):
                        scenarios.append(scenario)
                continue

            header_row, headers = find_header_row(worksheet)
            if not header_row:
                continue

            parsed_any_sheet = True
            columns = map_manual_columns(headers)
            if "steps" not in columns:
                continue

            sheet_name = worksheet.title
            for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                if not row or not any(
                    cell is not None and str(cell).strip() for cell in row
                ):
                    continue

                scenario = build_scenario_from_manual_row(
                    row,
                    columns,
                    sheet_name=sheet_name,
                )
                if scenario and scenario.get("functional_blocks"):
                    scenarios.append(scenario)

        if not parsed_any_sheet:
            raise ValueError("Manual test case header row not found in Excel.")
    finally:
        wb.close()

    return scenarios


def process_offline_generation(file_path):
    log(f"Offline mode: parsing manual Excel {file_path}")
    scenarios = parse_offline_manual_excel(file_path)

    # Optional module filter from config/app-input.json description
    try:
        import json
        from pathlib import Path as _P
        # services/engine/config/ — three levels above this file (uts_engine/exporters/xpedite/).
        _inp = _P(__file__).resolve().parent.parent.parent.parent / "config" / "app-input.json"
        if _inp.exists():
            _desc = (json.loads(_inp.read_text(encoding="utf-8-sig")).get("description") or "").strip()
            if _desc:
                try:
                    from uts_engine.discovery.module_filter import matches_module, parse_module_from_description
                    _mod = parse_module_from_description(_desc)
                except Exception:
                    _mod = _desc
                if _mod:
                    before = len(scenarios)
                    scenarios = [
                        s for s in scenarios
                        if matches_module(
                            " ".join([
                                str(s.get("description") or ""),
                                str(s.get("name") or ""),
                                str(s.get("component") or ""),
                                str(s.get("display_name") or ""),
                            ]),
                            _mod,
                        )
                    ]
                    log(f"Module filter '{_mod}': {before} -> {len(scenarios)} scenarios")
    except Exception as _e:
        log(f"Module filter skipped: {_e}")

    if not scenarios:
        raise ValueError("No automatable test cases found in the Excel file.")

    used_names = set()
    for scenario in scenarios:
        base_name = scenario.get("name", "Unnamed_Scenario")
        unique_name = base_name
        suffix = 2
        while unique_name in used_names:
            unique_name = f"{base_name}_{suffix}"
            suffix += 1
        used_names.add(unique_name)
        scenario["name"] = unique_name

    generate_manual_excel(scenarios)
    generate_shared_testdata_workbook(scenarios)

    for scenario in scenarios:
        normalize_scenario_steps(scenario)
        generate_suite(scenario)

    print(f"[OK] Offline: generated {len(scenarios)} Xpedite test suite(s) from Excel.")
    log(f"Offline SUCCESS: {len(scenarios)} scenario(s)")
    return scenarios


# -------------------------------------------------
# AI Call
# -------------------------------------------------
def interpret_user_prompt(user_prompt, doc_text=None, image_folder=None):

    requested_count = extract_requested_count(user_prompt)

    system_prompt = f"""
You are an automation framework designer for Xpedite.

CRITICAL:
You MUST return JSON in EXACT format:

{{
  "scenarios": [
    {{
      "name": "Scenario_Name",
      "functional_blocks": [
      {{
          "bc_name": "Launch_Application",
          "steps": [
            {{
              "object_name": "ApplicationURL",
              "action": "InvokeURL",
              "input_value": "<ApplicationURL>:https://opensource-demo.orangehrmlive.com",
              "expected_value": "",
              "description": "Launch application URL"
            }}
          ]
        }},
        {{
          "bc_name": "Block_Name",
          "steps": [
            {{
              "object_name": "",
              "action": "",
              "input_value": "",
              "expected_value": "",
              "description": ""
            }}
          ]
        }}
      ]
    }}
  ]
}}

GENERAL RULES:
1. Generate up to {requested_count} UNIQUE scenarios.
2. EVERY test case / scenario MUST include Launch_Application (InvokeURL) as its own BC, then a Login BC (Username + Password + Login click). Do not skip login.
3. Each scenario must represent a DIFFERENT business functionality after Launch + Login.
4. Do NOT create BC_Username or BC_Password alone — those steps belong in the Login BC. InvokeURL must remain its own BC (Launch_Application).
5. After Login, create BCs according to business functionality (one BC per module/screen flow).
6. After login, explore different modules or business areas of the application.
7. Cover a variety of realistic business workflows.

8. For data parameterisation, there are 2 types of variables:

   Static variable:
   - Denoted within angle brackets <>
   - Example: <Username>, <Password>, <AccountName>
   - These values come from datasheet and remain fixed during execution.

   Dynamic variable:
   - Denoted within square brackets []
   - Example: [OpportunityID], [QuoteNumber], [PortID]
   - These values are generated during execution and reused later.

9. You MUST use these variable formats when appropriate:
   - Use < > for input test data.
   - Use [ ] for system-generated or captured values.
   
10. Input Value Format Rule (MANDATORY):

For any step where action requires an input_value (example: SetText, SetPassword, PerformSelect):

The input_value MUST follow this format:

<StaticVariable>:DefaultValue

Example:

Correct:

input_value: "<Username>:admin_user"
input_value: "<Password>:Pass@123"
input_value: "<AccountName>:TestAccount"

If no static variable name is defined in the document, generate a meaningful variable and assign a default value.

Example:

Correct:

input_value: "<FirstName>:John"
input_value: "<LastName>:Doe"
input_value: "<Email>:john.doe@test.com"

Wrong:

input_value: "John"
input_value: "<Username>"
input_value: "admin"

11. If application generates an ID, store it using Dynamic variable format.

Example:
Correct:
input_value: "<OpportunityName>:TestOpportunity"

Correct:
expected_value: "[OpportunityID]"

Wrong:
input_value: "[OpportunityID]:12345"

DOCUMENT RULE:
12. The input document may contain:
   - Functional Specification Document (FSD) text
   - Screenshots (UI images)
   - Or a combination of both
13. Handling rules:
a. If FSD text is present:
   - Treat it as PRIMARY source of business flow
   - Follow steps exactly in sequence
b. If screenshots are present:
   - Extract visible UI elements (labels, buttons, fields) as object_name
   - Identify navigation between screens
   - Each screen transition MUST create a new functional_block
c. If both FSD text and screenshots are present:
   - Use FSD text for flow sequence
   - Use screenshots only to identify exact UI object names
   - Do NOT change flow based on images
14. Navigation Rule:
- Each page/screen/navigation change (from text or screenshot) MUST result in a new functional_block
- Do NOT merge multiple screens into one block
15. If document is unclear:
- Infer logical navigation but DO NOT invent new business flows
16. For screenshots:
    - Use images ONLY to extract exact visible object_name.
    - Do NOT assume that every image requires an action.
    - If image shows confirmation or summary screen, generate VerifyText instead of click.
    - If image shows read-only data, do not generate SetText.
17. If document text and image conflict, prioritize document description.

OBJECT RULES:
18. object_name MUST match EXACT visible UI label text.
19. Preserve exact casing as displayed in UI.
20. Do NOT include words like button, field, textbox, link, dropdown.
21. Do NOT invent technical IDs or internal names.

ACTION RULES:
22. Use ONLY:
   - InvokeURL
   - SetText
   - SetPassword
   - PerformSelect
   - PerformClick
   - VerifyText
   - VerifyNonEditableText
   
23. PerformClick MUST ALWAYS have:
   input_value = ""
   expected_value = ""

24. PerformClick is ONLY for clicking and never carries data.
25. For actions SetText, SetPassword, and PerformSelect:

input_value MUST always follow the format

<StaticVariable>:DefaultValue

Example:

Correct:

{{
  "object_name": "Username",
  "action": "SetText",
  "input_value": "<Username>:admin",
  "expected_value": "",
  "description": "Enter Username"
}}

VERIFY ACTION RULE (MANDATORY):

25. For ANY action that starts with "Verify"
   (VerifyNonEditableText)

   expected_value MUST ALWAYS be:
   "True"

26. input_value is expected result:
    The verification target value should go in input_value.
    input_value always contains a static variable written as:
    You MUST automatically convert it to:
 
    <VariableName>:DefaultValue
    
27. Never put actual expected result in expected_value. Always use "True".
   The verification target value should go in input_value if needed.

STRUCTURE RULES (MANDATORY):
28. BC form MUST match Xpedite UI (like screenshot):
    1) bc_name Launch_Application — ONE step: object ApplicationURL, action InvokeURL,
       input_value <ApplicationURL>:url
    2) bc_name Login — steps:
       - Username / SetText / <Username>:value
       - Password / SetPassword / <Password>:value
       - Sign in (or Login) / PerformClick
    3) Then dynamic business BCs (e.g. Dashboard) with VerifyNonEditableText rows:
       object NonEditableText, expected_value True, input_value <ExpectedResult>:text
    Example: BC_Launch_Application → BC_Login → BC_Dashboard → BC_<DynamicModule> …
29. Each block must have bc_name and steps.
30. Each scenario must represent a DIFFERENT business functionality after Launch_Application + Login.
31. Each step must contain:
    object_name, action, input_value,
    expected_value, description.
32. STATIC VARIABLE NORMALIZATION RULE (MANDATORY)

Before returning the final JSON, you MUST validate every step.

If any input_value contains a static variable written as:

<VariableName>

You MUST automatically convert it to:

<VariableName>:DefaultValue

Where DefaultValue is a meaningful sample value.

Examples:

<FirstName>      → <FirstName>:John
<LastName>       → <LastName>:Doe
<Username>       → <Username>:admin
<Password>       → <Password>:Pass@123
<Email>          → <Email>:test.user@example.com
<AccountName>    → <AccountName>:TestAccount

The final JSON MUST NEVER contain:

<VariableName>

All static variables MUST include :DefaultValue.

33. FINAL VALIDATION STEP (CRITICAL)

Before returning the JSON:

1. Scan every step.
2. If input_value matches pattern:

<VariableName>

3. Replace it with:

<VariableName>:DefaultValue

4. Only then return the JSON.
34. Return ONLY valid JSON.
35. Do NOT wrap output in markdown.

BC GRANULARITY RULES:

36. Create a SEPARATE functional_block for EACH application screen/page (logical business groups).
37. Do NOT combine multiple business screens into one functional_block.
38. If navigation moves to a new page, start a new functional_block.
39. Minimum 1 functional_block per major document section.
40. For long enterprise flows (like Salesforce CPQ), expect 4–8 functional_blocks.
41. Do NOT collapse entire journey into 1 or 2 blocks (Login may be shared; business blocks must still differ per scenario).
42. PerformClick MUST ALWAYS have empty input_value and empty expected_value.

INPUT VALUE STRICT RULE (CRITICAL)

43. For ANY action that requires input_value (SetText, SetPassword, PerformSelect):

The input_value MUST ALWAYS contain a default value.

MANDATORY FORMAT:

<StaticVariable>:DefaultValue

Examples:

Correct:
"<FirstName>:John"
"<LastName>:Doe"
"<Username>:admin"
"<AccountName>:TestAccount"

INCORRECT (DO NOT GENERATE):
"<FirstName>"
"<Username>"
"John"
"admin"

If input_value is generated without ":DefaultValue",
the response is INVALID and must be corrected.

44. NEVER generate a static variable without a default value.

The following pattern is STRICTLY FORBIDDEN:

<VariableName>

Correct format MUST ALWAYS be:

<VariableName>:DefaultValue

APPLICATION ENTRY RULE (MANDATORY)

44. The VERY FIRST step of the FIRST functional_block MUST ALWAYS launch the application using InvokeURL.

Requirements:

The first functional_block represents the application entry page.

The first step MUST be:

action: InvokeURL
object_name: ApplicationURL
input_value: "<ApplicationURL>:<url>"
expected_value: ""

Rules:

InvokeURL MUST appear only as the first step of the first functional_block.

No other action is allowed before InvokeURL.

If login is required, the Login functional_block must appear AFTER the InvokeURL block.

The InvokeURL step must not be skipped or placed later in the scenario.

45. If action is SetText, SetPassword, or PerformSelect:

input_value MUST ALWAYS follow:

<VariableName>:DefaultValue

There are NO exceptions.

ENTRY VALIDATION RULE

46. Before returning the JSON:

Verify that the first functional_block first step action = InvokeURL.

If missing, automatically insert a Launch_Application functional_block with InvokeURL as the first step.

47. TEXT AND OBJECT HANDLING RULE (CRITICAL)

- Preserve step text exactly as provided (no rephrasing, no space removal, no formatting changes).
- Always keep the full original sentence in description.

- For actionable steps (PerformClick, SetText, PerformSelect):
  - object_name MUST NOT be empty.
  - Derive object_name from the UI element or action target (e.g., "User Menu", "Validate", "Currency bought").

- If a step contains multiple values (e.g., "Currency bought, Currency sell"):
  - Split into multiple steps ONLY for execution.
  - Preserve original sentence in description of the first step.

- Ensure step numbering is sequential and corrected if duplicates exist.

MOBILE APPLICATION ENTRY RULE (MANDATORY)

If the user test case explicitly mentions:
- "Open Mobile application"
- "Launch Mobile App"
- "Open App"
- "Mobile application"

Then DO NOT generate InvokeURL.

Then the launch step MUST be generated as an application launch action.

Generate EXACTLY:

{{
  "bc_name": "Launch_Application",
  "steps": [
    {{
      "object_name": "Application Package",
      "action": "InvokeApplication",
      "input_value": "",
      "expected_value": "",
      "description": "Open Mobile application"
    }}
  ]
}}

Rules:
1. Use InvokeApplication ONLY for mobile application launch.
2. Use InvokeURL ONLY for web/browser applications.
3. Detect application type from user prompt.
4. If prompt contains URL → use InvokeURL.
5. If prompt contains Mobile application/App → use InvokeApplication.
6. Launch block must always remain the first functional_block.
7. Never generate InvokeURL for mobile flows.
8. Do not generate performClick action for mobile Application invoke step
"""

    combined_text = f"User Requirement:\n{user_prompt}\n"

    if re.search(r"orangehrm|opensource-demo\.orangehrmlive\.com", user_prompt, re.IGNORECASE):
        combined_text += """

Application-Specific Guidance:
- Application: OrangeHRM demo site
- URL: https://opensource-demo.orangehrmlive.com/
- Login credentials: Admin / admin123
- Generate detailed business workflows for Xpedite.
- Prefer realistic create, update, delete, and verify scenarios.
- Reuse a dedicated login functional block after the launch block.
- Prefer modules that support record maintenance, such as employee management or user management.
- Prioritize these scenario families:
  1. Employee create and verify in PIM.
  2. Employee update and verify in PIM.
  3. Employee delete/search verify in PIM.
  4. Admin user create/update/delete/verify in Admin when enough detail is needed.
- Each OrangeHRM scenario must contain explicit verification steps after create, after update, and after delete/search.
- Use realistic OrangeHRM labels such as Username, Password, Login, PIM, Add Employee, Employee List, Save, Search, Admin, Job Title, User Role, Employee Name, Status, Delete, Confirm Delete when relevant from flow.
- Prefer search-before-update and search-before-delete flows so the business workflow is execution-ready.
- Scenario names should be business specific, for example Employee_Create_And_Verify, Employee_Update_And_Verify, Employee_Delete_And_Verify.
"""

    if doc_text:
        combined_text += f"\nDocument Content:\n{doc_text}\n"

    return call_ai(
        system_prompt,
        combined_text,
        image_folder
    )

# -------------------------------------------------
# XML Builders
# -------------------------------------------------
def build_bc_xml(template_path, steps):
    tree = ET.parse(template_path)
    root = tree.getroot()

    for node in root.findall("Table"):
        root.remove(node)

    for step in steps:
        finalize_step(step)

        # Convert VerifyText to VerifyNonEditableText
        if step.get("action") == "VerifyText":
            step["action"] = "VerifyNonEditableText"
            step["object_name"] = "NonEditableText"
            
        table = ET.Element("Table")
        #ET.SubElement(table, "ObjectName").text = step.get("object_name", "")
        ET.SubElement(table, "ObjectName").text = (step.get("object_name", "") or "").replace(" Button", "").replace(" button", "")
        #input_value = normalize_input_value(step.get("input_value", ""))
        ET.SubElement(table, "InputValue").text = step.get("input_value", "")
        ET.SubElement(table, "ExpectedValue").text = step.get("expected_value", "")
        ET.SubElement(table, "Action").text = normalize_action(step.get("action"))
        ET.SubElement(table, "WindowType")
        ET.SubElement(table, "Step_Descriptions").text = step.get("description", "")
        ET.SubElement(table, "Remarks")
        ET.SubElement(table, "PreferenceSearch")
        ET.SubElement(table, "CreateReport").text = "true"
        ET.SubElement(table, "ImageID").text = "-1"
        root.append(table)

    return tree


def build_bpw_xml(template_path, bc_list):
    tree = ET.parse(template_path)
    root = tree.getroot()

    for node in root.findall("Table"):
        root.remove(node)

    for seq, bc in enumerate(bc_list, start=1):
        table = ET.Element("Table")
        ET.SubElement(table, "SeqNo").text = str(seq)
        ET.SubElement(table, "KeywordDrivenTestDataPath").text = bc
        ET.SubElement(table, "IterationCountForSameData").text = ""
        ET.SubElement(table, "IterationCountForDifferentData").text = ""
        ET.SubElement(table, "Skip").text = "false"
        ET.SubElement(table, "Remarks").text = ""
        ET.SubElement(table, "isLooped").text = "false"
        ET.SubElement(table, "BCVersion").text = "Latest"
        root.append(table)

    return tree


def collapse_one_line_blocks(blocks):
    """
    Required BC order:
    BC_InvokeURL → BC_Login → BC_<BusinessFunction> → ...
    """
    if not blocks:
        return blocks

    def _protected(name):
        n = (name or "").lower().replace(" ", "_")
        return n in ("launch_application", "invokeurl", "invoke_url", "login") or "launch" in n

    folded = []
    for block in blocks:
        steps = list(block.get("steps") or [])
        name = str(block.get("bc_name") or "Business_Flow")
        if not steps:
            continue
        if "launch" in name.lower() or name.lower() in ("invokeurl", "invoke_url"):
            name = "Launch_Application"
        if name.lower() in ("username", "password"):
            if folded and str(folded[-1].get("bc_name") or "").lower() == "login":
                folded[-1]["steps"].extend(steps)
            else:
                folded.append({"bc_name": "Login", "steps": steps})
            continue
        if "verify" in name.lower() and folded and "verify" in str(folded[-1].get("bc_name") or "").lower() and not _protected(str(folded[-1].get("bc_name") or "")):
            folded[-1]["steps"].extend(steps)
            continue
        folded.append({"bc_name": name, "steps": steps})

    result = []
    for block in folded:
        steps = list(block.get("steps") or [])
        name = str(block.get("bc_name") or "Business_Flow")
        if not steps:
            continue
        if _protected(name):
            result.append({"bc_name": name, "steps": steps})
            continue
        result.append({"bc_name": name, "steps": steps})

    return [
        {"bc_name": b.get("bc_name") or "Business_Flow", "steps": b["steps"]}
        for b in result
        if b.get("steps")
    ]


def ensure_login_block(scenario):
    """Every test case must have a Login BC (Username + Password + Login)."""
    blocks = scenario.get("functional_blocks") or []
    has_login = False
    for block in blocks:
        name = str(block.get("bc_name") or "").lower()
        actions = " ".join(
            f"{s.get('action','')} {s.get('object_name','')}" for s in (block.get("steps") or [])
        ).lower()
        if "login" in name or ("setpassword" in actions and "settext" in actions):
            has_login = True
            break
    if has_login:
        return scenario

    variables = parse_test_data_variables(scenario.get("test_data") or "")
    username = variables.get("<Username>", "test_user")
    password = variables.get("<Password>", "Pass@123")
    login_block = {
        "bc_name": "Login",
        "steps": [
            make_step("Username", "SetText", f"<Username>:{username}", "", "Enter username"),
            make_step("Password", "SetPassword", f"<Password>:{password}", "", "Enter password"),
            make_step("Login", "PerformClick", "", "", "Click Login"),
        ],
    }
    insert_at = 0
    for i, block in enumerate(blocks):
        if "launch" in str(block.get("bc_name") or "").lower():
            insert_at = i + 1
            break
    blocks = list(blocks)
    blocks.insert(insert_at, login_block)
    scenario["functional_blocks"] = blocks
    return scenario


def merge_functional_blocks(scenario):
    blocks = scenario.get("functional_blocks", [])
    scenario = ensure_login_block(scenario)
    blocks = scenario.get("functional_blocks", [])
    reworked = []
    login_bucket = []
    for block in blocks:
        name = str(block.get("bc_name") or "").lower()
        steps = block.get("steps") or []
        is_loginish = (
            name in ("username", "password", "login")
            or (len(steps) == 1 and str(steps[0].get("object_name") or "").lower()
                in ("username", "password", "login"))
        )
        if is_loginish:
            login_bucket.extend(steps)
            continue
        if login_bucket:
            reworked.append({"bc_name": "Login", "steps": login_bucket})
            login_bucket = []
        reworked.append(block)
    if login_bucket:
        reworked.append({"bc_name": "Login", "steps": login_bucket})
    scenario["functional_blocks"] = collapse_one_line_blocks(reworked)
    return scenario

# -------------------------------------------------
# Smart Scenario Parsing
# -------------------------------------------------
def extract_scenarios_from_ai(data):
    if "scenarios" in data:
        return data["scenarios"]

    if "functional_blocks" in data:
        return [data]

    if isinstance(data, list):
        return data

    return []


# -------------------------------------------------
# Generate test data file
# -------------------------------------------------
def generate_testdata_file(tc_folder, tc_name, variables):

    # File directly inside TC folder
    file_path = os.path.join(tc_folder, f"{tc_name}.tab")

    headers = [
        "Id", "ScenarioName", "MainBatchName", "Variable", "Value",
        "French", "Dutch", "Italian", "Remarks",
        "Project", "SubProject",
        "EnvDBQuery", "French_EnvDBQuery", "Dutch_EnvDBQuery", "Italian_EnvDBQuery",
        "DefaultOBR"
    ]

    with open(file_path, "w", encoding="utf-8") as f:

        # Header
        f.write("\t".join(headers) + "\n")

        for idx, (var, val) in enumerate(variables.items(), start=1):

            row = [
                str(idx),
                "TestCase",
                tc_name,
                var,
                f"\"{val}\"",
                "\"\"", "\"\"", "\"\"", "\"\"",
                "Demo",
                "Demo",
                "\"\"", "\"\"", "\"\"", "\"\"",
                "\"\""
            ]

            f.write("\t".join(row) + "\n")


def generate_testdata_workbook(tc_folder, tc_name, variables):
    Workbook, _, _ = import_openpyxl()
    file_path = os.path.abspath(os.path.join(tc_folder, f"{tc_name}.xlsx"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Keep the workbook as variable/value rows for direct Xpedite consumption.
    for var, val in variables.items():
        ws.append([var, val])

    if not variables:
        ws.append(["<SampleData>", "SampleValue"])

    wb.save(file_path)
    return file_path


def collect_scenario_data_variables(scenario):
    variables = dict(parse_test_data_variables(scenario.get("test_data", "")))
    for var, val in extract_static_variables(scenario).items():
        variables.setdefault(var, val)
    return variables


def generate_shared_testdata_workbook(scenarios):
    """Build output\\TestData.xlsx used by BC_GetTestData / BC_GetData."""
    template = os.path.join(
        TEMPLATE_PATH,
        "TestData.xlsx",
    )

    output = os.path.join(
        OUTPUT_PATH,
        "TestData.xlsx",
    )

    ExcelWriter(template).write(
        scenarios,
        output,
    )

    log(f"Shared test data written: {output} ({len(scenarios)} row(s))")
    return output


def generate_suite(scenario):
    base_name, tc_name, bpw_name = resolve_output_names(scenario)
    os.makedirs(OUTPUT_PATH, exist_ok=True)

    tc_template = resolve_template_path("TC_template.xml")
    bpw_template = resolve_template_path("BPW_template.xml")
    bc_template = resolve_template_path("BC_template.xml")
    tc_tree = ET.parse(tc_template)
    for elem in tc_tree.getroot().iter():
        if elem.text == "TC_HRM":
            elem.text = tc_name
        if elem.text == "BPW_HRM":
            elem.text = bpw_name

    tc_tree.write(os.path.join(OUTPUT_PATH, f"{tc_name}.xml"),
                  encoding="utf-8", xml_declaration=True)

    tc_folder = os.path.join(OUTPUT_PATH, tc_name)
    bpw_folder = os.path.join(tc_folder, bpw_name)
    os.makedirs(bpw_folder, exist_ok=True)

    # Remove legacy GetData BC files from older generator naming (BC_GetTestData_<TC>, etc.)
    for legacy_pattern in ("BC_GetTestData_*.xml", "BC_GetData_*.xml"):
        for legacy_file in glob.glob(os.path.join(bpw_folder, legacy_pattern)):
            try:
                os.remove(legacy_file)
            except OSError:
                pass
    
    # Extract variables
    variables = extract_static_variables(scenario)

    # Generate test data file
    generate_testdata_file(tc_folder, tc_name, variables)
    excel_path = generate_testdata_workbook(tc_folder, tc_name, variables)
    shared_excel_path = os.path.abspath(os.path.join(OUTPUT_PATH, SHARED_TESTDATA_FILE))

    bc_list = []

    # -----------------------------
    # First BC : GetData
    # -----------------------------
    getTestdata_bc = "BC_GetTestData"
    bc_list.append(getTestdata_bc)

    getdata_steps = [{
        "object_name": "[RowNum]",
        "input_value": f"{shared_excel_path};Sheet1;RowNum;AutomationTCName='+[@QcTcName]+';1;2",
        "expected_value": "",
        "action": "GetExcelRowValues",
        "description": "Get variable data from External file"
    }]

    bc_tree = build_bc_xml(bc_template, getdata_steps)
    bc_tree.write(
        os.path.join(bpw_folder, f"{getTestdata_bc}.xml"),
        encoding="utf-8",
        xml_declaration=True
    )

    # -----------------------------
    # Second BC : GetData
    # -----------------------------
    getdata_bc = "BC_GetData"
    bc_list.append(getdata_bc)

    getdata_steps = [{
        "object_name": f"{shared_excel_path};Sheet1;1;+[RowNum]+;row",
        "input_value": "True",
        "expected_value": "",
        "action": "GetDataFrmFile",
        "description": "Get variable data from External file"
    }]

    bc_tree = build_bc_xml(bc_template, getdata_steps)
    bc_tree.write(
        os.path.join(bpw_folder, f"{getdata_bc}.xml"),
        encoding="utf-8",
        xml_declaration=True
    )

    for index, block in enumerate(scenario.get("functional_blocks", []), start=1):
        bc_name = f"BC_{sanitize_name(block.get('bc_name') or f'Block_{index}')}"
        bc_list.append(bc_name)

        steps = block.get("steps", [])
        bc_tree = build_bc_xml(bc_template, steps)
        bc_tree.write(os.path.join(bpw_folder, f"{bc_name}.xml"),
                      encoding="utf-8", xml_declaration=True)

    log(f"BPW BC list for {tc_name}: {bc_list}")
    bpw_tree = build_bpw_xml(bpw_template, bc_list)
    bpw_tree.write(os.path.join(tc_folder, f"{bpw_name}.xml"),
                   encoding="utf-8", xml_declaration=True)

def process_generation(user_prompt, file_path=None, offline_mode=False):
    log(f"Prompt: {user_prompt}")
    log(f"File path: {file_path}")
    log(f"Offline mode: {offline_mode}")

    if file_path and file_path.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        if offline_mode or is_manual_testcase_workbook(file_path):
            return process_offline_generation(file_path)

    if offline_mode:
        raise ValueError(
            "Offline mode requires a manual test case Excel file with columns "
            "like TC ID and Test Steps."
        )

    doc_text = None
    image_folder = None

    if file_path:
        doc_text, image_folder = load_input_content(file_path)

    if not image_folder and folder_has_images(get_screenshots_folder()):
        image_folder = get_screenshots_folder()
        log(f"Using screenshots folder: {image_folder}")

    raw_response = interpret_user_prompt(user_prompt, doc_text, image_folder)
    cleaned_json = clean_json_response(raw_response)

    try:
        data = json.loads(cleaned_json)
    except json.JSONDecodeError:
        print("[ERROR] AI returned invalid JSON.")
        log("Invalid JSON from AI")
        print(raw_response)
        raise

    scenarios = extract_scenarios_from_ai(data)

    for scenario in scenarios:
        merge_functional_blocks(scenario)

    for scenario in scenarios:
        normalize_scenario_steps(scenario)

    generate_manual_excel(scenarios)
    generate_shared_testdata_workbook(scenarios)

    if not scenarios:
        print("[ERROR] No scenarios returned.")
        log("No scenarios returned")
        raise ValueError("No scenarios returned")

    for scenario in scenarios:
        generate_suite(scenario)

    print(f"[OK] Generated {len(scenarios)} test suite(s) successfully.")
    log("SUCCESS")
    return scenarios


def load_processed_files(state_file):
    if not os.path.exists(state_file):
        return {}

    try:
        with open(state_file, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            return data
    except Exception as exc:
        log(f"Unable to read watcher state: {exc}")

    return {}


def save_processed_files(state_file, processed_files):
    with open(state_file, "w", encoding="utf-8") as f:
        json.dump(processed_files, f, indent=2)


def watch_input_folder(watch_folder, user_prompt, poll_seconds=5, offline_mode=True):
    os.makedirs(watch_folder, exist_ok=True)
    state_file = os.path.join(watch_folder, ".generator_watch_state.json")
    processed_files = load_processed_files(state_file)

    print(f"Watching folder: {watch_folder}")
    log(f"Watching folder: {watch_folder}")

    while True:
        try:
            current_files = []
            for entry in os.scandir(watch_folder):
                if not entry.is_file():
                    continue
                if entry.name.startswith("."):
                    continue
                if not entry.name.lower().endswith(WATCH_EXTENSIONS):
                    continue
                current_files.append(entry.path)

            for file_path in sorted(current_files):
                modified_time = os.path.getmtime(file_path)
                modified_key = str(modified_time)

                if processed_files.get(file_path) == modified_key:
                    continue

                print(f"Processing file: {file_path}")
                log(f"Processing watched file: {file_path}")

                try:
                    process_generation(user_prompt, file_path, offline_mode=offline_mode)
                    processed_files[file_path] = modified_key
                    save_processed_files(state_file, processed_files)
                except Exception as exc:
                    log(f"Watch processing failed for {file_path}: {exc}")
                    traceback.print_exc()

            time.sleep(max(1, int(poll_seconds)))

        except KeyboardInterrupt:
            print("Watcher stopped by user.")
            log("Watcher stopped by user")
            break
        except Exception as exc:
            log(f"Watcher loop error: {exc}")
            traceback.print_exc()
            time.sleep(max(1, int(poll_seconds)))


def set_output_path(path):
    global OUTPUT_PATH
    if path:
        OUTPUT_PATH = path


# -------------------------------------------------
# MAIN
# -------------------------------------------------
if __name__ == "__main__":

    ensure_runtime_folders()
    log("-------------------------------------------------")
    log("Script started")

    watch_mode = False
    watch_folder = WATCH_DEFAULT_FOLDER
    poll_seconds = 5
    offline_mode = False

    # ✅ ALWAYS define first
    file_path = None

    log(f"sys.argv count: {len(sys.argv)}")
    log(f"sys.argv full: {sys.argv}")

    user_prompt = DEFAULT_PROMPT_TEXT
    all_args = [arg.strip() for arg in sys.argv[1:]]
    arg_index = 0

    if all_args and not all_args[0].startswith("--"):
        user_prompt = all_args[0]
        arg_index = 1

    while arg_index < len(all_args):
        arg = all_args[arg_index]

        if arg == "--watch":
            watch_mode = True

        elif arg == "--offline":
            offline_mode = True

        elif arg.startswith("--watch-folder="):
            watch_mode = True
            watch_folder = arg.split("=", 1)[1].strip() or WATCH_DEFAULT_FOLDER

        elif arg.startswith("--poll="):
            poll_seconds = int(arg.split("=", 1)[1].strip())

        elif arg.startswith("--output="):
            set_output_path(arg.split("=", 1)[1].strip())

        elif not arg.startswith("--") and not file_path:
            file_path = arg

        arg_index += 1

    if not file_path and DEFAULT_EXCEL_PATH:
        file_path = DEFAULT_EXCEL_PATH

    file_path = resolve_excel_path(file_path)
    if file_path:
        log(f"Excel file: {file_path}")
    else:
        log("No Excel file found in input folder")

    try:
        if watch_mode:
            if not user_prompt.strip() or user_prompt.startswith("--"):
                user_prompt = build_default_prompt()
            watch_input_folder(watch_folder, user_prompt, poll_seconds, offline_mode=True)
        else:
            if not file_path:
                print(f"ERROR: No Excel file found in: {INPUT_FOLDER}")
                print("Put your .xlsx file in the input\\ folder and run again.")
                log("ERROR: No Excel file in input folder")
                sys.exit(1)
            if not user_prompt.strip() or user_prompt.startswith("--"):
                user_prompt = build_default_prompt(file_path)
            if file_path and not offline_mode and file_path.lower().endswith(EXCEL_EXTENSIONS):
                offline_mode = is_manual_testcase_workbook(file_path)
            process_generation(user_prompt, file_path, offline_mode=offline_mode)

    except Exception as exc:
        log(f"ERROR: {exc}")
        traceback.print_exc()
        sys.exit(1)